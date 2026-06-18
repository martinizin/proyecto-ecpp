"""
Tests para ``ExportacionAsistenciaService`` (HU27).
"""

import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.asistencia.domain.services import AsistenciaCalculoService
from apps.asistencia.infrastructure.models import Asistencia
from apps.reportes.application.services import ExportacionAsistenciaService
from tests.factories import (
    AsistenciaFactory,
    MatriculaFactory,
    ParaleloFactory,
    PeriodoFactory,
)


pytestmark = pytest.mark.django_db


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _crear_paralelo_con_asistencia(periodo=None, num_estudiantes=2):
    """Crea un paralelo con N estudiantes y 10 sesiones de asistencia variadas."""
    periodo = periodo or PeriodoFactory(nombre="2026-A")
    paralelo = ParaleloFactory(periodo=periodo, nombre="A")
    for i in range(num_estudiantes):
        matricula = MatriculaFactory(paralelo=paralelo)
        # 8 presente, 1 ausente, 1 justificado = 10 sesiones, 90% asistencia
        for d in range(1, 11):
            estado = (
                Asistencia.Estado.PRESENTE
                if d < 9
                else (Asistencia.Estado.AUSENTE if d == 9 else Asistencia.Estado.JUSTIFICADO)
            )
            AsistenciaFactory(
                estudiante=matricula.estudiante,
                paralelo=paralelo,
                fecha=datetime.date(2026, 4, d),
                estado=estado,
            )
    return paralelo


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


class TestExportarExcelAsistencia:
    """Tests para ``ExportacionAsistenciaService.exportar_excel_asistencia``."""

    def test_exportar_excel_returns_bytesio(self, docente):
        """El método retorna un ``BytesIO``."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        result = service.exportar_excel_asistencia()
        assert isinstance(result, BytesIO)

    def test_exportar_excel_xlsx_magic_bytes(self, docente):
        """El archivo comienza con ``b\"PK\"``."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        assert buffer.getvalue()[:2] == b"PK"

    def test_exportar_excel_one_sheet_per_paralelo(self, docente):
        """Con 2 paralelos, el workbook tiene 2 sheets."""
        periodo = PeriodoFactory(nombre="2026-A")
        ParaleloFactory(periodo=periodo, nombre="A")
        ParaleloFactory(periodo=periodo, nombre="B")
        service = ExportacionAsistenciaService(filtros={"periodo_id": periodo.id}, usuario=docente)
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        assert len(wb.sheetnames) == 2

    def test_exportar_excel_column_order_r6(self, docente):
        """La fila de headers (row 5) tiene el orden R6.

        R6: ``Cédula, Nombres, Total Clases, Total Presentes, Total Ausentes,
        Total Justificadas, Porcentaje Asistencia, Estado``.
        """
        paralelo = _crear_paralelo_con_asistencia(num_estudiantes=1)
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        ws = wb.active
        header = [ws.cell(row=5, column=c).value for c in range(1, 9)]
        assert header == [
            "Cédula",
            "Nombres",
            "Total Clases",
            "Total Presentes",
            "Total Ausentes",
            "Total Justificadas",
            "Porcentaje Asistencia",
            "Estado",
        ]

    def test_exportar_excel_porcentaje_uses_calculo_service(self, docente):
        """El valor de la columna ``Porcentaje Asistencia`` viene del domain service."""
        paralelo = _crear_paralelo_con_asistencia(num_estudiantes=1)
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        ws = wb.active
        # data row 6: total=10, presentes=8, ausentes=1, justificadas=1 → 90.00%
        # Domain service formula: (presentes + justificadas) / total = (8+1)/10 = 90%
        cell_value = ws.cell(row=6, column=7).value
        expected = AsistenciaCalculoService().calcular_porcentaje_asistencia(9, 10)
        assert Decimal(str(cell_value)) == expected
        assert Decimal(str(cell_value)) == Decimal("90.00")

    def test_exportar_excel_estado_uses_evaluar_riesgo(self, docente):
        """El valor de ``Estado`` viene de ``evaluar_riesgo(porcentaje_inasistencia)``."""
        paralelo = _crear_paralelo_con_asistencia(num_estudiantes=1)
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        ws = wb.active
        estado_cell = ws.cell(row=6, column=8).value
        # inasistencia = 10% → "rojo" (UMBRAL_INASISTENCIA = 5%)
        expected = AsistenciaCalculoService().evaluar_riesgo(Decimal("10.00"))
        assert estado_cell == expected
        assert estado_cell == "rojo"

    def test_exportar_excel_header_block(self, docente):
        """Las primeras 3 filas son header block ECPPP, luego spacer, luego columnas."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        ws = wb.active
        assert "ECPPP" in str(ws.cell(row=1, column=1).value)
        assert "Paralelo" in str(ws.cell(row=2, column=1).value)
        assert "Generado:" in str(ws.cell(row=3, column=1).value)

    def test_exportar_excel_empty_filter_returns_valid_file(self, docente):
        """Filtros sin paralelos → archivo válido (magic bytes)."""
        periodo = PeriodoFactory(nombre="2026-A")
        service = ExportacionAsistenciaService(filtros={"periodo_id": periodo.id}, usuario=docente)
        buffer = service.exportar_excel_asistencia()
        assert buffer.getvalue()[:2] == b"PK"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


class TestExportarPdfAsistencia:
    """Tests para ``ExportacionAsistenciaService.exportar_pdf_asistencia``."""

    def test_exportar_pdf_returns_bytesio(self, docente):
        """El método retorna un ``BytesIO``."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        result = service.exportar_pdf_asistencia()
        assert isinstance(result, BytesIO)

    def test_exportar_pdf_magic_bytes(self, docente):
        """El archivo comienza con ``b\"%PDF\"``."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_pdf_asistencia()
        assert buffer.getvalue()[:4] == b"%PDF"

    def test_exportar_pdf_portrait_a4(self, docente):
        """El PDF es portrait A4 (width < height) — R7 spec."""
        paralelo = _crear_paralelo_con_asistencia()
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": paralelo.periodo_id}, usuario=docente
        )
        buffer = service.exportar_pdf_asistencia()
        import re

        text = buffer.getvalue().decode("latin-1", errors="ignore")
        match = re.search(
            r"/MediaBox\s*\[\s*(\d+\.?\d*)\s+(\d+\.?\d*)\s+(\d+\.?\d*)\s+(\d+\.?\d*)\s*\]",
            text,
        )
        assert match is not None, f"MediaBox no encontrado en PDF. Text head: {text[:300]!r}"
        x1, y1, x2, y2 = (float(g) for g in match.groups())
        width = x2 - x1
        height = y2 - y1
        # A4 portrait: 595x842
        assert width < height, f"portrait A4 esperado, got {width:.1f}x{height:.1f}"

    def test_exportar_pdf_one_table_per_paralelo(self, docente):
        """Con 2 paralelos, hay 2 ``/Type /Page`` en el PDF body."""
        periodo = PeriodoFactory(nombre="2026-A")
        ParaleloFactory(periodo=periodo, nombre="A")
        ParaleloFactory(periodo=periodo, nombre="B")
        service = ExportacionAsistenciaService(filtros={"periodo_id": periodo.id}, usuario=docente)
        buffer = service.exportar_pdf_asistencia()
        text = buffer.getvalue().decode("latin-1", errors="ignore")
        # Cada /Type /Page cuenta como un page object (excluyendo /Pages)
        page_refs = text.count("/Type /Page") - text.count("/Type /Pages")
        assert page_refs == 2

    def test_exportar_pdf_empty_filter_returns_valid_file(self, docente):
        """Filtros sin paralelos → PDF válido (magic bytes + EOF)."""
        periodo = PeriodoFactory(nombre="2026-A")
        service = ExportacionAsistenciaService(filtros={"periodo_id": periodo.id}, usuario=docente)
        buffer = service.exportar_pdf_asistencia()
        assert buffer.getvalue()[:4] == b"%PDF"
        assert b"%%EOF" in buffer.getvalue()[-1024:]


# ---------------------------------------------------------------------------
# Filtros
# ---------------------------------------------------------------------------


class TestFiltrosAsistencia:
    """Tests para los filtros del servicio de asistencia."""

    def test_filtro_por_paralelo(self, docente):
        """El filtro por paralelo limita los paralelos exportados a 1."""
        periodo = PeriodoFactory(nombre="2026-A")
        p1 = ParaleloFactory(periodo=periodo, nombre="A")
        ParaleloFactory(periodo=periodo, nombre="B")
        service = ExportacionAsistenciaService(
            filtros={"periodo_id": periodo.id, "paralelo_id": p1.id}, usuario=docente
        )
        buffer = service.exportar_excel_asistencia()
        wb = load_workbook(buffer)
        assert len(wb.sheetnames) == 1
