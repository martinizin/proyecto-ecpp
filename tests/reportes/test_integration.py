"""
Tests de integración E2E para exportación de reportes (HU27).

Cubre los escenarios cross-cutting que no pertenecen a un solo servicio
o vista: cobertura por rol, filtros combinables, filename R8 byte-locked,
service-reuse guard tests (R11).
"""

import re
from datetime import datetime
from datetime import timezone as dt_timezone
from unittest import mock

import pytest
from django.core.cache import cache
from django.urls import reverse

from apps.calificaciones.domain.services import CalificacionValidationService
from tests.factories import (
    CalificacionFactory,
    EvaluacionFactory,
    MatriculaFactory,
    ParaleloFactory,
    PeriodoFactory,
    RegistroCalificacionParaleloFactory,
)


pytestmark = pytest.mark.django_db


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _crear_periodo_con_paralelo(nombre_periodo="2026-A", nombre_paralelo="A"):
    periodo = PeriodoFactory(nombre=nombre_periodo)
    paralelo = ParaleloFactory(periodo=periodo, nombre=nombre_paralelo)
    RegistroCalificacionParaleloFactory(paralelo=paralelo)
    for tipo, peso in [
        ("parcial1", 30.0),
        ("parcial2_10h", 30.0),
        ("examen_final", 40.0),
    ]:
        from decimal import Decimal

        ev = EvaluacionFactory(paralelo=paralelo, tipo=tipo, peso=Decimal(str(peso)))
        matricula = MatriculaFactory(paralelo=paralelo)
        CalificacionFactory(evaluacion=ev, estudiante=matricula.estudiante, nota=Decimal("18.00"))
    return periodo, paralelo


def _clear_cache():
    cache.clear()


# ---------------------------------------------------------------------------
# Cobertura end-to-end por rol
# ---------------------------------------------------------------------------


class TestRoleCoverage:
    """E2E: cada rol permitido recibe 200; estudiante recibe 403."""

    @pytest.mark.parametrize(
        "client_fixture,endpoint,formato",
        [
            ("docente_client", "reportes:exportar_calificaciones", "excel"),
            ("docente_client", "reportes:exportar_calificaciones", "pdf"),
            ("inspector_client", "reportes:exportar_calificaciones", "excel"),
            ("secretaria_client", "reportes:exportar_calificaciones", "pdf"),
            ("docente_client", "reportes:exportar_asistencia", "excel"),
            ("inspector_client", "reportes:exportar_asistencia", "pdf"),
            ("secretaria_client", "reportes:exportar_asistencia", "excel"),
        ],
    )
    def test_allowed_role_gets_200(self, request, client_fixture, endpoint, formato):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        client = request.getfixturevalue(client_fixture)
        url = reverse(endpoint)
        response = client.get(url, {"periodo": periodo.id, "formato": formato})
        assert response.status_code == 200, f"{client_fixture} {endpoint} {formato}"

    @pytest.mark.parametrize(
        "client_fixture,endpoint",
        [
            ("estudiante_client", "reportes:exportar_calificaciones"),
            ("estudiante_client", "reportes:exportar_asistencia"),
        ],
    )
    def test_estudiante_is_forbidden_403(self, request, client_fixture, endpoint):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        client = request.getfixturevalue(client_fixture)
        url = reverse(endpoint)
        response = client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert response.status_code == 403

    def test_anonymous_redirects_or_403(self, client):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        url = reverse("reportes:exportar_calificaciones")
        response = client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert response.status_code in (302, 403)


# ---------------------------------------------------------------------------
# Empty filter result end-to-end
# ---------------------------------------------------------------------------


class TestEmptyResultE2E:
    """E2E: filtros sin matches → archivo válido con header block."""

    def test_calificaciones_empty_periodo_returns_200(self, docente_client):
        _clear_cache()
        periodo = PeriodoFactory(nombre="2026-A")  # sin paralelos
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert response.status_code == 200
        assert response.content[:2] == b"PK"

    def test_calificaciones_empty_periodo_pdf(self, docente_client):
        _clear_cache()
        periodo = PeriodoFactory(nombre="2026-A")
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "pdf"})
        assert response.status_code == 200
        assert response.content[:4] == b"%PDF"

    def test_asistencia_empty_periodo_returns_200(self, docente_client):
        _clear_cache()
        periodo = PeriodoFactory(nombre="2026-A")
        url = reverse("reportes:exportar_asistencia")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert response.status_code == 200
        assert response.content[:2] == b"PK"

    def test_calificaciones_empty_materia_filter(self, docente_client):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        # Filtro de materia que no existe
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "materia": 99999, "formato": "excel"}
        )
        assert response.status_code == 200
        assert response.content[:2] == b"PK"


# ---------------------------------------------------------------------------
# Filtros combinables
# ---------------------------------------------------------------------------


class TestFiltrosCombinables:
    """E2E: filtros periodo+materia+paralelo+estado."""

    def test_filtro_periodo_y_materia(self, docente_client):
        _clear_cache()
        from io import BytesIO

        from openpyxl import load_workbook
        from tests.factories import AsignaturaFactory

        periodo = PeriodoFactory(nombre="2026-A")
        m1 = AsignaturaFactory(codigo="M1")
        m2 = AsignaturaFactory(codigo="M2")
        ParaleloFactory(periodo=periodo, asignatura=m1, nombre="A")
        ParaleloFactory(periodo=periodo, asignatura=m2, nombre="B")
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "materia": m1.id, "formato": "excel"}
        )
        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        assert len(wb.sheetnames) == 1

    def test_filtro_periodo_y_paralelo(self, docente_client):
        _clear_cache()
        from io import BytesIO

        from openpyxl import load_workbook

        periodo, paralelo = _crear_periodo_con_paralelo()
        ParaleloFactory(periodo=periodo, nombre="B")
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "paralelo": paralelo.id, "formato": "excel"}
        )
        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        assert len(wb.sheetnames) == 1

    def test_filtro_estado_borrador(self, docente_client):
        _clear_cache()
        from decimal import Decimal
        from io import BytesIO

        from openpyxl import load_workbook
        from apps.calificaciones.infrastructure.models import RegistroCalificacionParalelo

        periodo = PeriodoFactory(nombre="2026-A")
        # Paralelo 1: estado borrador
        p1 = ParaleloFactory(periodo=periodo, nombre="A")
        RegistroCalificacionParaleloFactory(
            paralelo=p1, estado=RegistroCalificacionParalelo.Estado.BORRADOR
        )
        ev = EvaluacionFactory(paralelo=p1, tipo="parcial1", peso=Decimal("100.00"))
        m1 = MatriculaFactory(paralelo=p1)
        CalificacionFactory(evaluacion=ev, estudiante=m1.estudiante, nota=Decimal("18.00"))
        # Paralelo 2: estado validado
        p2 = ParaleloFactory(periodo=periodo, nombre="B")
        RegistroCalificacionParaleloFactory(
            paralelo=p2, estado=RegistroCalificacionParalelo.Estado.VALIDADO
        )

        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "estado": "borrador", "formato": "excel"}
        )
        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content))
        # Solo p1 (borrador) → 1 sheet
        assert len(wb.sheetnames) == 1

    def test_filtro_estado_invalido_devuelve_vacio(self, docente_client):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "estado": "foo", "formato": "excel"}
        )
        # estado inválido no matchea nada → 200 con archivo válido
        assert response.status_code == 200
        assert response.content[:2] == b"PK"


# ---------------------------------------------------------------------------
# Filename convention R8
# ---------------------------------------------------------------------------


class TestFilenameR8:
    """E2E: filename byte-locked según R8."""

    def test_filename_calificaciones_with_paralelo(self, docente_client):
        _clear_cache()
        periodo, paralelo = _crear_periodo_con_paralelo(
            nombre_periodo="2026A", nombre_paralelo="A"
        )
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(
            url, {"periodo": periodo.id, "paralelo": paralelo.id, "formato": "excel"}
        )
        match = re.search(r'filename="([^"]+)"', response["Content-Disposition"])
        assert match is not None
        filename = match.group(1)
        assert re.match(r"^calificaciones_2026A_A_\d{8}_\d{6}\.xlsx$", filename), filename

    def test_filename_calificaciones_without_paralelo(self, docente_client):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo(nombre_periodo="2026A")
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "excel"})
        match = re.search(r'filename="([^"]+)"', response["Content-Disposition"])
        assert match is not None
        filename = match.group(1)
        assert re.match(r"^calificaciones_2026A_\d{8}_\d{6}\.xlsx$", filename), filename

    def test_filename_asistencia_with_paralelo_pdf(self, docente_client):
        _clear_cache()
        periodo, paralelo = _crear_periodo_con_paralelo(
            nombre_periodo="2026A", nombre_paralelo="A"
        )
        url = reverse("reportes:exportar_asistencia")
        response = docente_client.get(
            url, {"periodo": periodo.id, "paralelo": paralelo.id, "formato": "pdf"}
        )
        match = re.search(r'filename="([^"]+)"', response["Content-Disposition"])
        assert match is not None
        filename = match.group(1)
        assert re.match(r"^asistencia_2026A_A_\d{8}_\d{6}\.pdf$", filename), filename

    def test_filename_asistencia_without_paralelo_pdf(self, docente_client):
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo(nombre_periodo="2026A")
        url = reverse("reportes:exportar_asistencia")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "pdf"})
        match = re.search(r'filename="([^"]+)"', response["Content-Disposition"])
        assert match is not None
        filename = match.group(1)
        assert re.match(r"^asistencia_2026A_\d{8}_\d{6}\.pdf$", filename), filename

    def test_archivo_es_descargable_no_renderizado(self, docente_client):
        """Content-Disposition es ``attachment``, no ``inline``."""
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        url = reverse("reportes:exportar_calificaciones")
        response = docente_client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert "attachment" in response["Content-Disposition"]
        assert "inline" not in response["Content-Disposition"]


# ---------------------------------------------------------------------------
# Service-reuse guard tests (R11)
# ---------------------------------------------------------------------------


class TestServiceReuseR11:
    """Guard tests: no fórmula re-implementada en los servicios de export."""

    def test_calificaciones_service_does_not_contain_promedio_formula(self):
        """``services.py`` no contiene la fórmula del promedio ponderado."""
        import inspect

        from apps.reportes import application

        source = inspect.getsource(application.services)
        # Buscamos patrones de la fórmula de promedio ponderado
        assert "* Decimal(str(peso))" not in source
        assert "/ total_peso" not in source
        assert "total_peso == 0" not in source

    def test_asistencia_service_does_not_contain_porcentaje_formula(self):
        """``services.py`` no contiene la fórmula de porcentaje de asistencia."""
        import inspect

        from apps.reportes import application

        source = inspect.getsource(application.services)
        # El cálculo real está en AsistenciaCalculoService
        # Verificamos que no haya un cálculo manual de porcentaje en reports
        assert "Decimal(sesiones_asistidas) / Decimal(total_sesiones)" not in source
        assert "* 100" not in source  # ratio × 100 es la firma del cálculo

    def test_calificaciones_service_calls_validation_service(self, docente, monkeypatch):
        """El servicio de calificaciones invoca el domain service de promedio."""
        from apps.reportes.application.services import ExportacionCalificacionesService

        call_count = {"n": 0}
        original = CalificacionValidationService.calcular_promedio_ponderado

        def counting(*args, **kwargs):
            call_count["n"] += 1
            return original(*args, **kwargs)

        monkeypatch.setattr(
            CalificacionValidationService, "calcular_promedio_ponderado", staticmethod(counting)
        )

        periodo, _ = _crear_periodo_con_paralelo()
        service = ExportacionCalificacionesService(
            filtros={"periodo_id": periodo.id}, usuario=docente
        )
        service.exportar_excel_calificaciones()
        assert call_count["n"] > 0, "calcular_promedio_ponderado no fue invocado"

    def test_asistencia_service_calls_calculo_service(self, docente, monkeypatch):
        """El servicio de asistencia invoca el domain service de cálculo."""
        from apps.asistencia.domain.services import AsistenciaCalculoService
        from apps.reportes.application.services import ExportacionAsistenciaService

        call_count = {"n": 0}
        original = AsistenciaCalculoService.calcular_porcentaje_asistencia

        def counting(self_, *args, **kwargs):
            call_count["n"] += 1
            return original(self_, *args, **kwargs)

        monkeypatch.setattr(AsistenciaCalculoService, "calcular_porcentaje_asistencia", counting)

        periodo, _ = _crear_periodo_con_paralelo()
        service = ExportacionAsistenciaService(filtros={"periodo_id": periodo.id}, usuario=docente)
        service.exportar_excel_asistencia()
        assert call_count["n"] > 0, "calcular_porcentaje_asistencia no fue invocado"


# ---------------------------------------------------------------------------
# Rate limit global cross-endpoint
# ---------------------------------------------------------------------------


class TestRateLimitGlobal:
    """R3: el rate limit es global entre calificaciones y asistencia."""

    # Instante fijo para el bucket del rate limiter. Congelamos ``timezone.now``
    # del limiter durante los requests para que el contador por minuto no se
    # resetee si el reloj cruza un borde de minuto en medio del test (CI lento).
    # El test y el limiter calculan así la MISMA cache key.
    _FIXED_NOW = datetime(2026, 1, 1, 12, 0, 0, tzinfo=dt_timezone.utc)
    _RATE_LIMITER_NOW = "apps.reportes.application.rate_limiter.timezone.now"

    def test_rate_limit_is_global_across_endpoints(self, docente_client):
        """10 calls (mezclando endpoints) pasan; la 11ª retorna 429."""
        _clear_cache()
        periodo, _ = _crear_periodo_con_paralelo()
        url_calif = reverse("reportes:exportar_calificaciones")
        url_asist = reverse("reportes:exportar_asistencia")

        with mock.patch(self._RATE_LIMITER_NOW, return_value=self._FIXED_NOW):
            # 5 calificaciones + 5 asistencia = 10
            for _ in range(5):
                response = docente_client.get(
                    url_calif, {"periodo": periodo.id, "formato": "excel"}
                )
                assert response.status_code == 200
            for _ in range(5):
                response = docente_client.get(
                    url_asist, {"periodo": periodo.id, "formato": "excel"}
                )
                assert response.status_code == 200

            # La 11ª (cualquier endpoint) → 429
            response = docente_client.get(url_calif, {"periodo": periodo.id, "formato": "excel"})
            assert response.status_code == 429

    def test_rate_limit_429_message_in_spanish(self, docente_client):
        """El body del 429 es JSON con shape ``{error, retry_after_seconds}`` (HU27b R18)."""
        _clear_cache()
        import json

        from django.contrib.auth import get_user_model

        Usuario = get_user_model()
        user = Usuario.objects.filter(rol="docente").order_by("-id").first()
        bucket = self._FIXED_NOW.strftime("%Y%m%d%H%M")
        cache.set(f"export_rate_{user.id}_{bucket}", 10, timeout=60)

        periodo, _ = _crear_periodo_con_paralelo()
        url = reverse("reportes:exportar_calificaciones")
        with mock.patch(self._RATE_LIMITER_NOW, return_value=self._FIXED_NOW):
            response = docente_client.get(url, {"periodo": periodo.id, "formato": "excel"})
        assert response.status_code == 429
        assert response["Content-Type"].startswith("application/json")
        body = json.loads(response.content)
        assert body == {"error": "rate_limit", "retry_after_seconds": 60}
